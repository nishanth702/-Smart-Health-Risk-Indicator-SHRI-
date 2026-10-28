import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc_path = r"C:\Users\nishanth golakoti\.gemini\antigravity\brain\bc073917-f8b5-475c-8814-f68446c85a19\.user_uploaded\media_1792932299600_46c733a9.docx"
doc = docx.Document(doc_path)

for i in range(125, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"[{i}] {p.style.name} | {p.text[:120]}")
