import docx

doc_path = r"C:\Users\nishanth golakoti\.gemini\antigravity\brain\bc073917-f8b5-475c-8814-f68446c85a19\.user_uploaded\media_1792932299600_46c733a9.docx"
doc = docx.Document(doc_path)

print("Paragraphs count:", len(doc.paragraphs))
print("Tables count:", len(doc.tables))
print("Sections count:", len(doc.sections))

for s_idx, sec in enumerate(doc.sections):
    print(f"\nSection {s_idx}:")
    print("  Margins (inches): Top:", sec.top_margin.inches, "Bottom:", sec.bottom_margin.inches, "Left:", sec.left_margin.inches, "Right:", sec.right_margin.inches)
    print("  Page Width:", sec.page_width.inches, "Height:", sec.page_height.inches)

print("\nFirst 40 non-empty paragraphs:")
count = 0
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.style.name} | {p.text[:100]}")
        count += 1
        if count >= 40:
            break
